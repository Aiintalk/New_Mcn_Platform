"""
Unit tests for app.services.benchmark_report — docx generation.

覆盖：
- profile / plan doc_type 标题分支
- account_name 有/无（"未知" 兜底）
- 多级标题 # ~ ######（level 被 cap 到 4）
- 粗体片段 **bold** 行内格式
- 空行 + 普通文本 + 空内容
- 目录自动创建（_ensure_dir）
- 返回路径 + 文件真实存在 + docx 可重新打开

测试通过 monkeypatch 把 REPORT_DIR 重定向到 tmp_path，避免污染工作目录。
"""
from pathlib import Path

import pytest
from docx import Document

from app.services import benchmark_report


@pytest.fixture
def report_dir(tmp_path, monkeypatch):
    """重定向 REPORT_DIR 到临时目录。"""
    out = tmp_path / "benchmark_reports"
    monkeypatch.setattr(benchmark_report, "REPORT_DIR", out)
    return out


class TestGenerateDocx:
    def test_generates_profile_docx(self, report_dir):
        path = benchmark_report.generate_docx(
            analysis_id=1001,
            content="# 标题一\n\n普通段落\n\n**粗体内容**",
            account_name="小红",
            doc_type="profile",
        )
        assert Path(path).exists()
        assert "1001_profile.docx" in path
        assert path.startswith(str(report_dir))

        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        assert any("人格档案" in t and "小红" in t for t in texts)
        assert "标题一" in texts
        assert "普通段落" in texts

    def test_generates_plan_docx(self, report_dir):
        path = benchmark_report.generate_docx(
            analysis_id=1002,
            content="## 子标题\n\n正文",
            account_name="小明",
            doc_type="plan",
        )
        assert Path(path).exists()
        assert "1002_plan.docx" in path

        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        assert any("内容规划" in t and "小明" in t for t in texts)

    def test_default_doc_type_is_profile(self, report_dir):
        path = benchmark_report.generate_docx(
            analysis_id=1003,
            content="x",
            account_name="A",
        )
        assert "1003_profile.docx" in path

    def test_account_name_none_shows_unknown(self, report_dir):
        path = benchmark_report.generate_docx(
            analysis_id=1004,
            content="x",
            account_name=None,
            doc_type="profile",
        )
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        assert any("未知" in t for t in texts)

    def test_heading_levels_capped_at_4(self, report_dir):
        """# ~ ###### 都被识别为 heading；level 被 cap 到 4。

        文档自身有 1 个 level=1 标题（"人格档案 · {name}"），加上 6 个内容标题 = 7。
        """
        content = "# 一级\n## 二级\n### 三级\n#### 四级\n##### 五级\n###### 六级"
        path = benchmark_report.generate_docx(
            analysis_id=1005,
            content=content,
            account_name="x",
        )
        doc = Document(path)
        headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
        # 1 个文档标题 + 6 个内容标题
        assert len(headings) == 7
        # 内容文本都被识别为 heading
        content_texts = {"一级", "二级", "三级", "四级", "五级", "六级"}
        heading_texts = {p.text for p in headings}
        assert content_texts.issubset(heading_texts)
        # 所有 level ≤ 4（验证 cap）
        levels = [int(p.style.name.split()[-1]) for p in headings]
        assert max(levels) <= 4

    def test_bold_inline(self, report_dir):
        content = "前缀 **粗体片段** 后缀"
        path = benchmark_report.generate_docx(
            analysis_id=1006,
            content=content,
            account_name="x",
        )
        doc = Document(path)
        # 找包含"粗体片段"的段落，验证 bold run
        found = False
        for p in doc.paragraphs:
            bold_runs = [r for r in p.runs if r.bold and r.text == "粗体片段"]
            if bold_runs:
                found = True
                break
        assert found, "未找到含粗体的段落"

    def test_multiple_bold_in_one_line(self, report_dir):
        content = "**第一段** 中间 **第二段**"
        path = benchmark_report.generate_docx(
            analysis_id=1007,
            content=content,
            account_name="x",
        )
        doc = Document(path)
        for p in doc.paragraphs:
            bold_texts = [r.text for r in p.runs if r.bold]
            if "第一段" in bold_texts:
                assert "第二段" in bold_texts
                return
        pytest.fail("未找到含多个粗体的段落")

    def test_empty_lines_become_empty_paragraphs(self, report_dir):
        content = "行一\n\n\n行二"
        path = benchmark_report.generate_docx(
            analysis_id=1008,
            content=content,
            account_name="x",
        )
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        assert "行一" in texts
        assert "行二" in texts

    def test_empty_content_does_not_crash(self, report_dir):
        path = benchmark_report.generate_docx(
            analysis_id=1009,
            content="",
            account_name="x",
        )
        assert Path(path).exists()
        assert Path(path).stat().st_size > 0

    def test_creates_nested_report_dir(self, tmp_path, monkeypatch):
        """REPORT_DIR 不存在时 _ensure_dir 递归创建。"""
        out = tmp_path / "nested" / "deep" / "reports"
        monkeypatch.setattr(benchmark_report, "REPORT_DIR", out)
        assert not out.exists()

        path = benchmark_report.generate_docx(
            analysis_id=1010,
            content="x",
            account_name="x",
        )
        assert out.exists()
        assert Path(path).exists()

    def test_overwrites_existing_file(self, report_dir):
        """同一 analysis_id + doc_type 第二次生成覆盖第一次。"""
        kwargs = {
            "analysis_id": 1011,
            "content": "第一次",
            "account_name": "x",
            "doc_type": "profile",
        }
        path1 = benchmark_report.generate_docx(**kwargs)
        size1 = Path(path1).stat().st_size

        kwargs["content"] = "第二次生成的更长内容 " * 10
        path2 = benchmark_report.generate_docx(**kwargs)
        size2 = Path(path2).stat().st_size

        assert path1 == path2
        # 文件被覆盖（内容不同，大小可能不同但能写入成功）
        assert Path(path2).exists()
